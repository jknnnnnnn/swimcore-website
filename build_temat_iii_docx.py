from pathlib import Path
import math
import statistics

import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "Temat_III_rozwiazania.docx"
CHART_PATH = OUT_DIR / "Temat_III_zadanie_2_wykres.png"


def regularized_gamma_p(a, x):
    if x <= 0:
        return 0.0
    gln = math.lgamma(a)
    if x < a + 1.0:
        ap = a
        total = 1.0 / a
        delta = total
        for _ in range(1000):
            ap += 1.0
            delta *= x / ap
            total += delta
            if abs(delta) < abs(total) * 1e-14:
                break
        return total * math.exp(-x + a * math.log(x) - gln)

    b = x + 1.0 - a
    c = 1e30
    d = 1.0 / b
    h = d
    for i in range(1, 1000):
        an = -i * (i - a)
        b += 2.0
        d = an * d + b
        if abs(d) < 1e-30:
            d = 1e-30
        c = b + an / c
        if abs(c) < 1e-30:
            c = 1e-30
        d = 1.0 / d
        delta = d * c
        h *= delta
        if abs(delta - 1.0) < 1e-14:
            break
    q = math.exp(-x + a * math.log(x) - gln) * h
    return 1.0 - q


def chi2_cdf(x, df):
    return regularized_gamma_p(df / 2.0, x / 2.0)


def inv_chi2(p, df):
    lo = 0.0
    hi = max(float(df), 1.0)
    while chi2_cdf(hi, df) < p:
        hi *= 2.0
    for _ in range(160):
        mid = (lo + hi) / 2.0
        if chi2_cdf(mid, df) < p:
            lo = mid
        else:
            hi = mid
    return (lo + hi) / 2.0


def simpson(f, a, b):
    c = (a + b) / 2.0
    return (b - a) * (f(a) + 4.0 * f(c) + f(b)) / 6.0


def adaptive_simpson(f, a, b, eps, whole, depth):
    c = (a + b) / 2.0
    left = simpson(f, a, c)
    right = simpson(f, c, b)
    if depth <= 0 or abs(left + right - whole) <= 15.0 * eps:
        return left + right + (left + right - whole) / 15.0
    return (
        adaptive_simpson(f, a, c, eps / 2.0, left, depth - 1)
        + adaptive_simpson(f, c, b, eps / 2.0, right, depth - 1)
    )


def integrate(f, a, b, eps=1e-12):
    if a == b:
        return 0.0
    if a > b:
        return -integrate(f, b, a, eps)
    return adaptive_simpson(f, a, b, eps, simpson(f, a, b), 40)


def make_t_dist(df):
    coeff = math.gamma((df + 1.0) / 2.0) / (
        math.sqrt(df * math.pi) * math.gamma(df / 2.0)
    )

    def pdf(z):
        return coeff * (1.0 + z * z / df) ** (-(df + 1.0) / 2.0)

    def cdf(z):
        if z == 0:
            return 0.5
        area = integrate(pdf, 0.0, abs(z))
        return 0.5 + area if z > 0 else 0.5 - area

    def inv(p):
        lo, hi = -20.0, 20.0
        for _ in range(120):
            mid = (lo + hi) / 2.0
            if cdf(mid) < p:
                lo = mid
            else:
                hi = mid
        return (lo + hi) / 2.0

    return pdf, cdf, inv


def fmt(value, places=3):
    return f"{value:.{places}f}".replace(".", ",")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="A6A6A6", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths_inches):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths_dxa = [int(round(width * 1440)) for width in widths_inches]
    total_dxa = sum(widths_dxa)

    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = tbl.tblGrid
    if tbl_grid is not None:
        tbl.remove(tbl_grid)
    tbl_grid = OxmlElement("w:tblGrid")
    for width_dxa in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width_dxa))
        tbl_grid.append(grid_col)
    tbl.insert(1, tbl_grid)

    for row in table.rows:
        for idx, width_dxa in enumerate(widths_dxa):
            cell = row.cells[idx]
            cell.width = Inches(width_dxa / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width_dxa))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")


def add_table(doc, headers, rows, widths, center_cols=None):
    center_cols = set(center_cols or [])
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    hdr_cells = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr_cells[i].text = text
        set_cell_shading(hdr_cells[i], "F2F4F7")
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            cells[i].paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            )
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


def add_h2(doc, text):
    paragraph = doc.add_paragraph(style="Heading 2")
    paragraph.add_run(text)
    return paragraph


def add_body(doc, text="", bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)
    return paragraph


def add_formula(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    return paragraph


def get_font(size=22, bold=False):
    candidates = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def draw_scatter_chart(x_values, y_values, r_value):
    width, height = 1200, 760
    left, top, right, bottom = 130, 80, 70, 120
    plot_w = width - left - right
    plot_h = height - top - bottom
    x_min, x_max = 10, 75
    y_min, y_max = 1, 10

    def sx(x):
        return left + (x - x_min) / (x_max - x_min) * plot_w

    def sy(y):
        return top + (y_max - y) / (y_max - y_min) * plot_h

    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = get_font(24)
    small = get_font(20)
    title_font = get_font(34, bold=True)

    draw.text((left, 25), "Zależność wartości produkcji od liczby zatrudnionych", fill=(30, 30, 30), font=title_font)
    draw.rectangle([left, top, width - right, height - bottom], outline=(80, 80, 80), width=2)

    for tick in range(10, 80, 10):
        px = sx(tick)
        draw.line([px, top, px, height - bottom], fill=(226, 226, 226), width=1)
        label = str(tick)
        bbox = draw.textbbox((0, 0), label, font=small)
        draw.text((px - (bbox[2] - bbox[0]) / 2, height - bottom + 12), label, fill=(60, 60, 60), font=small)
    for tick in range(1, 11):
        py = sy(tick)
        draw.line([left, py, width - right, py], fill=(226, 226, 226), width=1)
        label = str(tick)
        bbox = draw.textbbox((0, 0), label, font=small)
        draw.text((left - 20 - (bbox[2] - bbox[0]), py - 12), label, fill=(60, 60, 60), font=small)

    slope, intercept = np.polyfit(x_values, y_values, 1)
    x_line = [x_min, x_max]
    y_line = [slope * x_min + intercept, slope * x_max + intercept]
    draw.line([sx(x_line[0]), sy(y_line[0]), sx(x_line[1]), sy(y_line[1])], fill=(31, 78, 121), width=4)

    for x, y in zip(x_values, y_values):
        px, py = sx(x), sy(y)
        draw.ellipse([px - 9, py - 9, px + 9, py + 9], fill=(46, 116, 181), outline=(10, 45, 80), width=2)

    x_label = "x_i - liczba zatrudnionych (osoby)"
    bbox = draw.textbbox((0, 0), x_label, font=font)
    draw.text((left + (plot_w - (bbox[2] - bbox[0])) / 2, height - 58), x_label, fill=(30, 30, 30), font=font)

    y_label = "y_i - wartość produkcji (mln zł)"
    y_img = Image.new("RGBA", (520, 40), (255, 255, 255, 0))
    y_draw = ImageDraw.Draw(y_img)
    y_draw.text((0, 0), y_label, fill=(30, 30, 30), font=font)
    y_img = y_img.rotate(90, expand=True)
    img.paste(y_img, (22, top + 25), y_img)

    note = f"r = {r_value:.3f}".replace(".", ",")
    draw.rounded_rectangle([width - 250, top + 20, width - 90, top + 72], radius=8, fill=(242, 244, 247), outline=(190, 190, 190))
    draw.text((width - 225, top + 34), note, fill=(30, 30, 30), font=font)
    img.save(CHART_PATH)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10


def build_document():
    noise = np.array([46, 48, 45, 42, 41, 50, 49, 44, 39, 51, 52, 38, 45], dtype=float)
    n = len(noise)
    df = n - 1
    mean = float(noise.mean())
    s = float(noise.std(ddof=1))
    s2 = s * s
    se = s / math.sqrt(n)
    alpha = 0.10
    chi2_005 = inv_chi2(alpha / 2.0, df)
    chi2_095 = inv_chi2(1.0 - alpha / 2.0, df)
    ci_var_low = df * s2 / chi2_095
    ci_var_high = df * s2 / chi2_005
    _, t_cdf, t_inv = make_t_dist(df)
    t_stat = (mean - 45.0) / se
    t_crit = t_inv(1.0 - alpha)
    p_value = 1.0 - t_cdf(t_stat)

    y = np.array([2, 4, 4, 6, 7, 7, 8, 9], dtype=float)
    x = np.array([15, 23, 30, 54, 44, 50, 56, 70], dtype=float)
    r = float(np.corrcoef(x, y)[0, 1])
    draw_scatter_chart(x, y, r)

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("TEMAT III - rozwiązania")
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Statystyka opisowa, estymacja, test hipotezy, korelacja i rozkład Poissona")
    run.italic = True
    run.font.color.rgb = RGBColor(85, 85, 85)

    doc.add_paragraph(style="Heading 1").add_run("Zadanie 1")
    add_body(
        doc,
        "Dane: 46, 48, 45, 42, 41, 50, 49, 44, 39, 51, 52, 38, 45 dB. "
        "Przyjmujemy, że jest to próba reprezentatywna z rozkładu normalnego."
    )

    add_h2(doc, "a) Średni poziom hałasu i odchylenie standardowe")
    add_formula(doc, "x̄ = Σx_i / n = 590 / 13 = 45,385 dB")
    add_formula(doc, "s² = Σ(x_i - x̄)² / (n - 1) = 20,423;   s = 4,519 dB")
    add_table(
        doc,
        ["Wielkość", "Wartość"],
        [
            ["Liczebność próby n", str(n)],
            ["Średnia z próby x̄", f"{fmt(mean)} dB"],
            ["Wariancja z próby s²", fmt(s2)],
            ["Odchylenie standardowe z próby s", f"{fmt(s)} dB"],
        ],
        [3.6, 2.9],
        center_cols=[1],
    )
    add_body(
        doc,
        "Interpretacja: przeciętny zmierzony poziom hałasu w próbie wynosi około "
        f"{fmt(mean)} dB, a typowe odchylenie pojedynczych pomiarów od średniej to około {fmt(s)} dB."
    )

    add_h2(doc, "b) Estymacja punktowa średniego poziomu hałasu")
    add_formula(doc, "μ̂ = x̄ = 45,385 dB")
    add_formula(doc, "SE(x̄) = s / √n = 4,519 / √13 = 1,253 dB")
    add_body(
        doc,
        "Punktowym oszacowaniem przeciętnego poziomu hałasu jest średnia z próby: "
        f"{fmt(mean)} dB. Średni błąd oszacowania wynosi {fmt(se)} dB."
    )

    add_h2(doc, "c) Przedział ufności dla wariancji przy 1 - α = 0,90")
    add_formula(doc, "df = n - 1 = 12;   χ²_0,05;12 = 5,226;   χ²_0,95;12 = 21,026")
    add_formula(
        doc,
        "P( (n-1)s² / χ²_0,95 ≤ σ² ≤ (n-1)s² / χ²_0,05 ) = 0,90",
    )
    add_formula(
        doc,
        f"σ² ∈ ({fmt(ci_var_low)}, {fmt(ci_var_high)}) dB²",
    )
    add_body(
        doc,
        "Interpretacja: z ufnością 90% wariancja poziomu hałasu w hali znajduje się w przedziale "
        f"od {fmt(ci_var_low)} do {fmt(ci_var_high)} dB²."
    )

    add_h2(doc, "d) Weryfikacja hipotezy, że średni poziom hałasu nie przekracza 45 dB")
    add_body(doc, "Przyjmujemy test jednostronny dla średniej przy nieznanym odchyleniu standardowym:")
    add_formula(doc, "H₀: μ ≤ 45 dB;   H₁: μ > 45 dB;   α = 0,10")
    add_formula(doc, "t = (x̄ - μ₀) / (s / √n) = (45,385 - 45) / 1,253 = 0,307")
    add_table(
        doc,
        ["Element testu", "Wartość"],
        [
            ["Liczba stopni swobody", str(df)],
            ["Statystyka testowa t", fmt(t_stat)],
            ["Wartość krytyczna t_0,90;12", fmt(t_crit)],
            ["p-value", fmt(p_value)],
            ["Decyzja", "brak podstaw do odrzucenia H₀"],
        ],
        [3.6, 2.9],
        center_cols=[1],
    )
    add_body(
        doc,
        f"Ponieważ t = {fmt(t_stat)} < {fmt(t_crit)} oraz p-value = {fmt(p_value)} > 0,10, "
        "nie odrzucamy hipotezy zerowej. Na poziomie istotności 0,10 nie ma podstaw, aby stwierdzić, "
        "że przeciętny poziom hałasu w hali przekracza 45 dB."
    )

    add_h2(doc, "e) Interpretacja wyników")
    add_body(
        doc,
        "Wyniki sugerują, że średni poziom hałasu jest bardzo bliski granicy 45 dB. "
        "Próba nie daje jednak statystycznie istotnego dowodu przekroczenia tej granicy. "
        "Zmienność pomiarów jest umiarkowana, co potwierdza odchylenie standardowe około 4,52 dB "
        "oraz dość szeroki przedział ufności dla wariancji."
    )

    doc.add_page_break()
    doc.add_paragraph(style="Heading 1").add_run("Zadanie 2")
    add_body(
        doc,
        "Dane opisują wartość produkcji y_i w mln zł oraz liczbę zatrudnionych x_i w osobach."
    )
    add_table(
        doc,
        ["i", "x_i (osoby)", "y_i (mln zł)"],
        [[str(i + 1), str(int(xv)), str(int(yv))] for i, (xv, yv) in enumerate(zip(x, y))],
        [0.8, 2.85, 2.85],
        center_cols=[0, 1, 2],
    )

    add_h2(doc, "a) Wykres zależności")
    chart_par = doc.add_paragraph()
    chart_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chart_par.add_run().add_picture(str(CHART_PATH), width=Inches(6.2))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    caption.add_run("Wykres punktowy z linią trendu dla par (x_i, y_i).").italic = True

    add_h2(doc, "b) Współczynnik korelacji")
    add_formula(
        doc,
        "r = Σ[(x_i - x̄)(y_i - ȳ)] / √(Σ(x_i - x̄)² · Σ(y_i - ȳ)²)",
    )
    add_formula(doc, f"r = {fmt(r)};   r² = {fmt(r * r)}")
    add_body(
        doc,
        f"Współczynnik korelacji Pearsona wynosi {fmt(r)}. Oznacza to bardzo silną dodatnią zależność liniową: "
        "w badanej próbie przedsiębiorstwa zatrudniające więcej osób osiągają zwykle większą wartość produkcji. "
        f"Wartość r² = {fmt(r * r)} wskazuje, że około {fmt(100 * r * r, 1)}% zróżnicowania wartości produkcji "
        "można opisać liniową zależnością od liczby zatrudnionych."
    )

    doc.add_paragraph(style="Heading 1").add_run("Zadanie 3")
    add_body(doc, "Rozkład rzadkich zdarzeń, czyli rozkład Poissona, opisuje liczbę wystąpień zdarzenia w ustalonym przedziale czasu, przestrzeni lub innej jednostce obserwacji, jeżeli zdarzenia pojawiają się niezależnie i ze stałą średnią intensywnością.")
    add_formula(doc, "P(X = k) = (λ^k · e^(-λ)) / k!,   k = 0, 1, 2, ...;   λ > 0")
    add_table(
        doc,
        ["Własność", "Opis"],
        [
            ["Parametr", "λ oznacza średnią liczbę zdarzeń w rozpatrywanym przedziale."],
            ["Wartość oczekiwana", "E(X) = λ."],
            ["Wariancja", "D²(X) = Var(X) = λ, a odchylenie standardowe wynosi √λ."],
            ["Niezależność przyrostów", "Liczby zdarzeń w rozłącznych przedziałach są niezależne."],
            ["Addytywność", "Suma niezależnych zmiennych Poissona ma rozkład Poissona z parametrem równym sumie parametrów."],
            ["Zastosowania", "Liczba awarii, wypadków, zgłoszeń, błędów drukarskich lub klientów w jednostce czasu."],
            ["Przybliżenie dwumianowego", "Dla dużego n i małego p rozkład B(n, p) można przybliżyć rozkładem Poissona z λ = np."],
        ],
        [1.8, 4.7],
        center_cols=[],
    )
    add_body(
        doc,
        "Interpretacja parametru jest prosta: jeżeli λ = 3, to w danym przedziale oczekujemy średnio trzech zdarzeń. "
        "Im większe λ, tym rozkład staje się mniej skośny i bardziej zbliża się kształtem do rozkładu normalnego."
    )

    doc.core_properties.title = "Temat III - rozwiązania"
    doc.core_properties.subject = "Statystyka"
    doc.core_properties.author = "Codex"
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = build_document()
    print(path)
