from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


OUT = Path(__file__).resolve().parent / "Temat_III_1do1_zdjecie.docx"


def set_run_font(run, size=13, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_margins(cell, top=80, start=95, bottom=80, end=95):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "12")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_table_geometry(table, widths_cm):
    widths_dxa = [int(round(width / 2.54 * 1440)) for width in widths_cm]
    total_dxa = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_dxa))

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "95")

    existing_grid = tbl.tblGrid
    if existing_grid is not None:
        tbl.remove(existing_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for idx, width in enumerate(widths_dxa):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def add_plain_paragraph(doc, text="", size=13, bold=False, before=0, after=0, line_spacing=1.0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def add_runs(paragraph, pieces, size=13):
    for text, opts in pieces:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=opts.get("bold", False))
        run.font.subscript = opts.get("subscript", False)


def add_task_item(doc, label, text, continuation=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.78)
    paragraph.paragraph_format.first_line_indent = Cm(-0.78)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.78), WD_TAB_ALIGNMENT.LEFT)
    add_runs(paragraph, [(label, {"bold": True}), ("\t" + text, {"bold": True})], size=13)
    if continuation:
        cont = doc.add_paragraph()
        cont.paragraph_format.left_indent = Cm(0.78)
        cont.paragraph_format.space_after = Pt(2)
        cont.paragraph_format.line_spacing = 1.0
        run = cont.add_run(continuation)
        set_run_font(run, size=13, bold=True)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.05)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(2.62)
    section.right_margin = Cm(1.28)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    add_plain_paragraph(doc, "T E M A T  III", size=13.5, bold=True, before=0, after=31)
    add_plain_paragraph(doc, "Zadanie 1", size=13.5, bold=True, after=22)

    add_plain_paragraph(
        doc,
        "W pewnej dużej hali produkcyjnej przeprowadzono badanie poziomu hałasu w różnych jej",
        bold=True,
    )
    add_plain_paragraph(
        doc,
        "miejscach i otrzymano następujące wyniki [xᵢ – w decybelach (dB)]:",
        bold=True,
    )
    add_plain_paragraph(doc, "46, 48, 45, 42, 41, 50, 49, 44, 39, 51, 52, 38, 45", bold=True)
    add_plain_paragraph(
        doc,
        "Traktując wyniki pomiarów jako próbę reprezentatywną hałasu na tej sali oraz zakładając że",
        bold=True,
    )
    add_plain_paragraph(doc, "poziom hałasu podlega rozkładowi normalnemu:", bold=True)

    add_task_item(
        doc,
        "a)",
        "obliczyć średni poziom hałasu oraz jego odchylenie standardowe w otrzymanej próbie,",
    )
    add_task_item(
        doc,
        "b)",
        "metodą estymacji punktowej oszacować przeciętny poziom hałasu na tej sali, podając",
        "średni błąd oszacowania,",
    )
    add_task_item(
        doc,
        "c)",
        "zbudować przedział ufności dla wariancji poziomu hałasu w tej hali, przyjmując",
        "współczynnik ufności 1 – α = 0,90,",
    )
    add_task_item(
        doc,
        "d)",
        "zweryfikować hipotezę o tym, że przeciętny poziom hałasu w tej hali nie przekracza 45",
        "dB, przyjmując poziom istotności α = 0,1,",
    )
    add_task_item(doc, "e)", "otrzymane wyniki każdorazowo zinterpretować.")

    add_plain_paragraph(doc, "Zadanie 2", size=13.5, bold=True, before=23)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_runs(
        p,
        [
            ("Wartość produkcji (y", {"bold": True}),
            ("i", {"bold": True, "subscript": True}),
            (" – w mln zł) oraz liczba zatrudnionych (x", {"bold": True}),
            ("i", {"bold": True, "subscript": True}),
            (" - w osobach) w", {"bold": True}),
        ],
        size=13,
    )
    add_plain_paragraph(doc, "przedsiębiorstwach pewnej branży były następujące:", bold=True)

    table = doc.add_table(rows=2, cols=9)
    set_table_borders(table)
    set_table_geometry(table, [3.0, 1.55, 1.55, 1.55, 1.55, 1.55, 1.55, 1.55, 1.55])
    table.rows[0].height = Cm(1.62)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    table.rows[1].height = Cm(1.84)
    table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    labels = [
        [("y", {"bold": True}), ("i", {"bold": True, "subscript": True}), (" (mln zł)", {"bold": True})],
        [("x", {"bold": True}), ("i", {"bold": True, "subscript": True}), (" (osoby)", {"bold": True})],
    ]
    values = [["2", "4", "4", "6", "7", "7", "8", "9"], ["15", "23", "30", "54", "44", "50", "56", "70"]]
    for r_idx, label_runs in enumerate(labels):
        p = table.cell(r_idx, 0).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_runs(p, label_runs, size=13)
        for c_idx, value in enumerate(values[r_idx], start=1):
            cell = table.cell(r_idx, c_idx)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, size=13)

    add_task_item(
        doc,
        "a)",
        "sporządzić wykres zależności pomiędzy tymi zmiennymi;",
    )
    add_task_item(
        doc,
        "b)",
        "obliczyć i zinterpretować współczynnik korelacji pomiędzy nimi.",
    )

    add_plain_paragraph(doc, "Zadanie 3", size=13.5, bold=True, before=71, after=5)
    add_plain_paragraph(doc, "Opisać rozkład rzadkich zdarzeń (Poissona) i jego własności.", size=13.5)

    doc.core_properties.title = "Temat III"
    doc.core_properties.subject = "Treść z fotografii"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
